import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any
from io import BytesIO

def create_students_excel(students: List[Dict[str, Any]]) -> BytesIO:
    """
    O'quvchilar ro'yxatini zamonaviy Excel fayl shaklida hosil qiladi.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lidlar Bazasi"

    # Sarlavhalar
    headers = [
        "№", "F.I.O", "Telefon Raqam", "Telegram ID", "Username",
        "Reklama Manbasi", "Bilim Darajasi", "Test Bali", 
        "Tanlangan Guruh Vaqti", "Kupon Kodi", "Status", 
        "Sinov Darsi", "Menejer Izohi", "Ro'yxatdan O'tgan Vaqti"
    ]
    ws.append(headers)

    # Dizayn stillari
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border_style = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Ma'lumotlarni yozish
    for idx, s in enumerate(students, 1):
        row_data = [
            idx,
            s.get("full_name") or "Noma'lum",
            s.get("phone") or "-",
            s.get("telegram_id") or "-",
            f"@{s.get('username')}" if s.get("username") else "-",
            s.get("utm_source") or "organic",
            s.get("level") or "-",
            s.get("test_score") or 0,
            s.get("preferred_schedule") or "-",
            s.get("coupon_code") or "-",
            s.get("status") or "Yangi",
            s.get("trial_date") or "-",
            s.get("manager_note") or "-",
            str(s.get("registered_at") or "")
        ]
        ws.append(row_data)

        row_idx = idx + 1
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border_style
            cell.alignment = left_align if col_idx in [2, 5, 9, 13] else center_align

    # Ustun kengliklarini avtomatik to'g'rilash
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
