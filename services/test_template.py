import re
import json
from io import BytesIO
from typing import List, Dict, Any, Tuple
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = Path(__file__).resolve().parent.parent

def generate_test_template_excel() -> BytesIO:
    """
    O'qituvchilar va administratorlar uchun Ingliz tili testi shablonini hosil qiladi.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Savollari"

    # Sarlavhalar
    headers = [
        "№",
        "Savol matni (Question)",
        "Variant A",
        "Variant B",
        "Variant C",
        "Variant D",
        "To'g'ri javob (A, B, C yoki D)",
        "Daraja (Level)"
    ]
    ws.append(headers)

    # Namunaviy savollar
    sample_rows = [
        [1, "She _______ to the gym every morning before breakfast.", "go", "goes", "is going", "went", "B", "Beginner"],
        [2, "We haven't seen them _______ last summer.", "for", "since", "during", "while", "B", "Elementary"],
        [3, "If I _______ you, I would take that opportunity immediately.", "am", "was", "were", "had been", "C", "Pre-Intermediate"],
        [4, "Despite _______ tired, he continued preparing for his IELTS exam.", "being", "be", "was", "to be", "A", "Intermediate"],
        [5, "By the time the teacher entered, the students _______ their assignment.", "finished", "have finished", "had finished", "finish", "C", "Intermediate"],
        [6, "Hardly _______ entered the room when the phone started ringing.", "he had", "had he", "he did", "did he", "B", "Upper-Intermediate"]
    ]

    for row in sample_rows:
        ws.append(row)

    # Dizayn stillari
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    border_style = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Sarlavhani bezash
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border_style

    # Namunaviy qatorlarni bezash
    for row_idx in range(2, len(sample_rows) + 2):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border_style
            if col_idx in [1, 7, 8]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Ustun kengliklari
    col_widths = {
        "A": 6,
        "B": 48,
        "C": 20,
        "D": 20,
        "E": 20,
        "F": 20,
        "G": 28,
        "H": 22
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 28

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def parse_test_excel(file_bytes: bytes) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Yuklangan Excel faylni o'qib, savollar ro'yxatiga aylantiradi.
    (questions, errors) qaytaradi.
    """
    stream = BytesIO(file_bytes)
    wb = openpyxl.load_workbook(stream, data_only=True)
    ws = wb.active

    questions: List[Dict[str, Any]] = []
    errors: List[str] = []

    answer_map = {
        "A": 0, "0": 0, "1": 0,
        "B": 1, "2": 1,
        "C": 2, "3": 2,
        "D": 3, "4": 3
    }

    # 2-qatordan boshlab o'qish (1-qator sarlavha)
    current_id = 1
    for row_idx in range(2, ws.max_row + 1):
        q_text = ws.cell(row=row_idx, column=2).value
        # Agar savol matni bo'sh bo'lsa, qatorni tashlab ketamiz
        if not q_text or not str(q_text).strip():
            continue

        q_text = str(q_text).strip()
        optA = str(ws.cell(row=row_idx, column=3).value or "").strip()
        optB = str(ws.cell(row=row_idx, column=4).value or "").strip()
        optC = str(ws.cell(row=row_idx, column=5).value or "").strip()
        optD = str(ws.cell(row=row_idx, column=6).value or "").strip()
        raw_ans = str(ws.cell(row=row_idx, column=7).value or "").strip().upper()
        level = str(ws.cell(row=row_idx, column=8).value or "Intermediate").strip()

        if not (optA and optB and optC and optD):
            errors.append(f"{row_idx}-qator: Variantlardan biri to'ldirilmagan!")
            continue

        if raw_ans not in answer_map:
            errors.append(f"{row_idx}-qator: To'g'ri javob noto'g'ri ko'rsatilgan ('{raw_ans}'). Faqat A, B, C yoki D bo'lishi kerak!")
            continue

        ans_idx = answer_map[raw_ans]

        # Standart daraja nomlari
        if "begin" in level.lower():
            level = "Beginner"
        elif "elem" in level.lower():
            level = "Elementary"
        elif "pre" in level.lower():
            level = "Pre-Intermediate"
        elif "upper" in level.lower():
            level = "Upper-Intermediate"
        elif "inter" in level.lower():
            level = "Intermediate"
        elif "adv" in level.lower():
            level = "Advanced"

        questions.append({
            "id": current_id,
            "question": f"{current_id}. {q_text}" if not re.match(r"^\d+\.", q_text) else q_text,
            "options": [optA, optB, optC, optD],
            "answer": ans_idx,
            "level": level
        })
        current_id += 1

    return questions, errors


def sync_questions_to_html(questions: List[Dict[str, Any]]):
    """
    Yangi savollarni web/templates/index.html va netlify_app/index.html fayllariga avtomatik yozadi.
    """
    json_str = json.dumps(questions, ensure_ascii=False, indent=6)
    replacement = f"const testQuestions = {json_str};"

    target_files = [
        BASE_DIR / "netlify_app" / "index.html",
        BASE_DIR / "web" / "templates" / "index.html"
    ]

    pattern = re.compile(r"const testQuestions = \[[\s\S]*?\];", re.MULTILINE)

    for fpath in target_files:
        if fpath.exists():
            try:
                content = fpath.read_text(encoding="utf-8")
                if pattern.search(content):
                    new_content = pattern.sub(replacement, content, count=1)
                    fpath.write_text(new_content, encoding="utf-8")
            except Exception as e:
                print(f"Error syncing questions to {fpath}: {e}")


def sync_courses_to_html(courses: List[Dict[str, Any]]):
    """
    Yangi kurslar ro'yxatini web/templates/index.html va netlify_app/index.html fayllariga yozadi.
    """
    json_str = json.dumps(courses, ensure_ascii=False, indent=6)
    replacement = f"const coursesData = {json_str};"

    target_files = [
        BASE_DIR / "netlify_app" / "index.html",
        BASE_DIR / "web" / "templates" / "index.html"
    ]

    pattern = re.compile(r"const coursesData = \[[\s\S]*?\];", re.MULTILINE)

    for fpath in target_files:
        if fpath.exists():
            try:
                content = fpath.read_text(encoding="utf-8")
                if pattern.search(content):
                    new_content = pattern.sub(replacement, content, count=1)
                    fpath.write_text(new_content, encoding="utf-8")
            except Exception as e:
                print(f"Error syncing courses to {fpath}: {e}")


def sync_results_to_html(results: List[Dict[str, Any]]):
    """
    Yangi o'quvchi natijalarini web/templates/index.html va netlify_app/index.html fayllariga yozadi.
    """
    json_str = json.dumps(results, ensure_ascii=False, indent=6)
    replacement = f"const resultsData = {json_str};"

    target_files = [
        BASE_DIR / "netlify_app" / "index.html",
        BASE_DIR / "web" / "templates" / "index.html"
    ]

    pattern = re.compile(r"const resultsData = \[[\s\S]*?\];", re.MULTILINE)

    for fpath in target_files:
        if fpath.exists():
            try:
                content = fpath.read_text(encoding="utf-8")
                if pattern.search(content):
                    new_content = pattern.sub(replacement, content, count=1)
                    fpath.write_text(new_content, encoding="utf-8")
            except Exception as e:
                print(f"Error syncing results to {fpath}: {e}")


def sync_schedules_to_html(schedules: List[Dict[str, Any]]):
    """
    Dars jadvallarini web/templates/index.html va netlify_app/index.html fayllariga yozadi.
    """
    json_str = json.dumps(schedules, ensure_ascii=False, indent=6)
    replacement = f"const schedulesData = {json_str};"

    target_files = [
        BASE_DIR / "netlify_app" / "index.html",
        BASE_DIR / "web" / "templates" / "index.html"
    ]

    pattern = re.compile(r"const schedulesData = \[[\s\S]*?\];", re.MULTILINE)

    for fpath in target_files:
        if fpath.exists():
            try:
                content = fpath.read_text(encoding="utf-8")
                if pattern.search(content):
                    new_content = pattern.sub(replacement, content, count=1)
                    fpath.write_text(new_content, encoding="utf-8")
            except Exception as e:
                print(f"Error syncing schedules to {fpath}: {e}")


